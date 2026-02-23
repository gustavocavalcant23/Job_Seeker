import os
import re
from datetime import datetime
from openpyxl import Workbook, load_workbook


class ReportService:

    @staticmethod
    def export_to_excel(query: str, jobs: list):
        today = datetime.now().strftime('%d-%m')

        formatted_query = re.sub(
            r'[^a-zA-Z0-9]',
            '',
            query.title().replace(' ', '')
        )

        filename = f'{formatted_query}{today}.xlsx'

        folder = 'job_reports'
        os.makedirs(folder, exist_ok=True)

        filepath = os.path.join(folder, filename)

        if os.path.exists(filepath):
            wb = load_workbook(filepath)
            ws = wb.active
            ws.delete_rows(2, ws.max_row)
        else: 
            wb = Workbook()
            ws = wb.active
            ws.title = 'Vagas'
            ws.append(['ID', 'Título', 'Descrição', 'URL', 'Origem'])

        for job in jobs:
            ws.append([
                job.id,
                job.title,
                job.description,
                job.url,
                job.origin,
            ])

        wb.save(filepath)
        print(f'Relatório salvo em: {filepath}')